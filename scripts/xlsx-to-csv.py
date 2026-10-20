#!/usr/bin/env python3
"""Dump every data tab of data/FundED-data.xlsx to data/sheet/<Tab>.csv, the files
`npm run sync:local` reads. Run this after editing the workbook (needs `openpyxl`).

Rows are written with every cell quoted, trailing blank columns trimmed to the header,
and dates as YYYY-MM-DD, so a round trip through build-sheet-template.py is a no-op."""
import csv
import sys
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "FundED-data.xlsx"
OUT = ROOT / "data" / "sheet"
TABS = ["Districts", "Tracks", "Stages", "Schools", "Needs", "Milestones", "Engage", "Inquiries", "Funding", "Sources", "Legend"]


def cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def main() -> int:
    if not SRC.exists():
        print(f"missing {SRC.relative_to(ROOT)}", file=sys.stderr)
        return 1
    wb = load_workbook(SRC, read_only=True, data_only=True)
    missing = [t for t in TABS if t not in wb.sheetnames]
    if missing:
        print(f"workbook is missing tabs: {', '.join(missing)}", file=sys.stderr)
        return 1
    for tab in TABS:
        ws = wb[tab]
        rows = [[cell_text(c) for c in r] for r in ws.iter_rows(values_only=True)]
        if not rows:
            print(f"{tab}: empty tab", file=sys.stderr)
            return 1
        header = rows[0]
        width = len(header)
        while width > 0 and header[width - 1] == "":
            width -= 1
        out_rows = []
        for r in rows:
            r = (r + [""] * width)[:width]
            if any(x != "" for x in r):
                out_rows.append(r)
        with open(OUT / f"{tab}.csv", "w", newline="", encoding="utf-8") as f:
            csv.writer(f, quoting=csv.QUOTE_ALL, lineterminator="\n").writerows(out_rows)
        print(f"{tab}: {len(out_rows) - 1} rows")
    return 0


if __name__ == "__main__":
    sys.exit(main())
