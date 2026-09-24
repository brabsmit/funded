#!/usr/bin/env python3
"""Build data/FundED-data-template.xlsx from data/sheet/*.csv with header styling,
frozen header row, and dropdown validation on enum columns. Upload the xlsx to
Google Drive and open with Google Sheets; validation carries over."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "sheet"
OUT = ROOT / "data" / "FundED-data-template.xlsx"
TABS = ["Tracks", "Stages", "Schools", "Needs", "Milestones", "Engage", "Inquiries", "Funding", "Sources", "Legend"]
ENUMS = {
    "basis": "fact,requirement,estimate,interpretation",
    "stage_basis": "fact,requirement,estimate,interpretation",
    "cost_basis": "fact,requirement,estimate,interpretation",
    "window_basis": "fact,requirement,estimate,interpretation",
    "why_track_basis": "fact,requirement,estimate,interpretation",
    "verification": "draft,documented,confirmed",
}
STATUS = {"Schools": "live,draft", "Milestones": "done,next,later", "Inquiries": "open,answered,partial,unanswered"}
FONT = "Arial"

# Cross-sheet reference validation: column name -> tab holding the id it points to.
# "every column ending in source_id" is handled separately below.
REF_TARGETS = {
    "track_id": "Tracks",
    "current_stage_id": "Stages",
    "school_id": "Schools",
    "need_id": "Needs",
    "parent_id": "Funding",
}
# Each tab's own id column is skipped (e.g. track_id on Tracks is not a reference).
OWN_ID_COLUMN = {
    "Tracks": "track_id", "Stages": "stage_id", "Schools": "school_id", "Needs": "need_id",
    "Milestones": "milestone_id", "Engage": "engage_id", "Inquiries": "inquiry_id", "Funding": "funding_id", "Sources": "source_id",
}

wb = Workbook()
wb.remove(wb.active)
ref_validation_counts = {}
for tab in TABS:
    ws = wb.create_sheet(tab)
    with open(SRC / f"{tab}.csv", newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    for r in rows:
        ws.append(r)
    header = rows[0]
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = PatternFill("solid", fgColor="E3EEF7")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT)
    ws.freeze_panes = "A2"
    ref_count = 0
    for i, name in enumerate(header, start=1):
        col = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col].width = max(14, min(60, max(len(str(r[i-1])) if i-1 < len(r) else 0 for r in rows) + 2))
        values = ENUMS.get(name) or (STATUS.get(tab) if name == "status" else None)
        if values and tab != "Legend":
            dv = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True, showErrorMessage=True,
                                errorTitle="Not an allowed value", error=f"Use one of: {values.replace(',', ', ')}")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}500")
        is_own_id = name == OWN_ID_COLUMN.get(tab)
        target = None if is_own_id else (REF_TARGETS.get(name) or ("Sources" if name.endswith("source_id") else None))
        if target and tab != "Legend":
            ref_dv = DataValidation(type="list", formula1=f"={target}!$A$2:$A$500", allow_blank=True, showErrorMessage=True,
                                errorTitle="Unknown reference", error=f"Must match an id already listed on {target}")
            ws.add_data_validation(ref_dv)
            ref_dv.add(f"{col}2:{col}500")
            ref_count += 1
    if ref_count:
        ref_validation_counts[tab] = ref_count
wb.save(OUT)
print(f"wrote {OUT.relative_to(ROOT)} with tabs: {', '.join(TABS)}")
print("cross-sheet reference validations per tab:")
for tab in TABS:
    print(f"  {tab}: {ref_validation_counts.get(tab, 0)}")
